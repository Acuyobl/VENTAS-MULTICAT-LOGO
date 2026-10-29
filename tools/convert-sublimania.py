#!/usr/bin/env python3
"""Convierte la hoja de pedidos de SUBLIMANIA CR en un catálogo CSV importable.

- Toma la primera hoja (pedidos) y la normaliza a columnas de catálogo.
- Extrae las fotos incrustadas de la hoja «Imagenes» y las embebe como data URI
  base64 en la columna Imagen, para que el importador las suba a MinIO.
- Precio sugerido = costo × 1.5 (el archivo no trae precio de venta).

Uso:
  python tools/convert-sublimania.py "C:\\Users\\...\\Pedido_...xlsx"
"""
import base64
import csv
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path(sys.argv[1])
OUT_DIR = Path("catalogs")
OUT_CSV = OUT_DIR / "05-sublimania-cr.csv"
ORIGINAL = OUT_DIR / "05-sublimania-cr-pedido-original.xlsx"

OUT_DIR.mkdir(exist_ok=True)
shutil.copy2(SOURCE, ORIGINAL)

book = load_workbook(SOURCE, data_only=True)
sheet = book["Pedido SUBLIMANIA CR"]
headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
index = {header: position for position, header in enumerate(headers)}

rows = []
for row in sheet.iter_rows(min_row=2, values_only=True):
    if not row or all(value is None for value in row) or str(row[0] or "").strip() == "":
        continue
    if str(row[index.get("ESTADO", len(row) - 1)] or "").strip().lower() in {"total", ""}:
        continue
    rows.append(row)

# --- imágenes incrustadas, asociadas por cliente ---
images_by_client = {}
image_sheet = book["Imagenes"]
for image in getattr(image_sheet, "_images", []):
    anchor_row = image.anchor._from.row + 1
    client = image_sheet.cell(row=anchor_row, column=2).value
    if not client:
        continue
    data = image._data()
    kind = "png" if data[:8] == b"\x89PNG\r\n\x1a\n" else "jpeg"
    if data[:2] != b"\xff\xd8" and kind == "jpeg":
        kind = "png" if data[:8] == b"\x89PNG\r\n\x1a\n" else "jpeg"
    images_by_client[str(client).strip().upper()] = f"data:image/{kind};base64,{base64.b64encode(data).decode()}"

print(f"filas de pedido: {len(rows)} · imágenes: {len(images_by_client)}")


def value(row, column):
    position = index.get(column)
    if position is None or position >= len(row) or row[position] is None:
        return ""
    return str(row[position]).strip()


out_rows = []
for number, row in enumerate(rows, start=1):
    client = value(row, "CLIENTE")
    name = value(row, "CAMISA") or value(row, "DETALLE") or f"Producto {number}"
    detail = value(row, "DETALLE")
    fabric = value(row, "TELA")
    kind = value(row, "TIPO")
    neck = value(row, "CUELLO")
    gender = value(row, "GENERO")
    sleeve = value(row, "MANGA")
    size = value(row, "TALLA")
    urgency = value(row, "URGENCIA")
    cost = int(float(value(row, "COSTO").replace(",", "") or 0))
    stock = int(float(value(row, "CANTIDAD").replace(",", "") or 1))
    price = int(round((cost * 1.5) / 100.0)) * 100

    description = " ".join(
        part
        for part in [
            f"{name}.",
            f"Detalle: {detail}." if detail else "",
            f"Tela {fabric}." if fabric else "",
            f"Tipo {kind}, cuello {neck.lower()}, manga {sleeve.lower()}." if kind else "",
            f"Talla {size}." if size else "",
            f"Urgencia {urgency.lower()}." if urgency else "",
        ]
        if part
    ).strip()
    image = images_by_client.get(client.upper(), "")

    out_rows.append(
        {
            "SKU": f"SUB-{number:03d}",
            "Nombre": name,
            "Descripción": description,
            "Precio": price,
            "Costo": cost,
            "Stock": stock,
            "Imagen": image,
            "Talla": size,
            "Tela": fabric,
            "Tipo": kind,
            "Cuello": neck,
            "Genero": gender,
            "Manga": sleeve,
            "Detalle": detail,
            "Urgencia": urgency,
        }
    )

columns = list(out_rows[0].keys())
with OUT_CSV.open("w", encoding="utf-8", newline="") as handle:
    writer = csv.DictWriter(handle, fieldnames=columns)
    writer.writeheader()
    writer.writerows(out_rows)

size_kb = OUT_CSV.stat().st_size / 1024
print(f"escrito {OUT_CSV} · {len(out_rows)} productos · {size_kb:.1f} KB · copia original {ORIGINAL.name}")
for item in out_rows:
    print(f"  {item['SKU']} {item['Nombre']} · ${item['Precio']} · talla {item['Talla']} · imagen={'sí' if item['Imagen'] else 'no'}")
