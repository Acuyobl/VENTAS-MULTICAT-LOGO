#!/usr/bin/env python3
"""Inspecciona un catálogo XLSX real: hojas, encabezados, combinadas e imágenes."""
import sys
from pathlib import Path

from openpyxl import load_workbook

path = Path(sys.argv[1])
print(f"archivo: {path} · {path.stat().st_size/1024:.1f} KB")
book = load_workbook(path, data_only=True)
print("hojas:", book.sheetnames)

for sheet in book.worksheets:
    print(f"\n=== hoja «{sheet.title}» · {sheet.max_row} filas × {sheet.max_column} columnas ===")
    merged = list(sheet.merged_cells.ranges)
    print(f"combinadas: {len(merged)}", [str(r) for r in merged[:15]])
    images = getattr(sheet, "_images", [])
    print(f"imágenes: {len(images)}")
    for img in images[:12]:
        try:
            anchor = img.anchor
            row = anchor._from.row + 1
            col = anchor._from.col + 1
            print(f"  imagen en fila {row}, columna {col} ({img.width}x{img.height})")
        except Exception as error:  # noqa: BLE001
            print("  imagen sin ancla:", error)

    max_rows = min(sheet.max_row, int(sys.argv[2]) if len(sys.argv) > 2 else 30)
    for row in sheet.iter_rows(min_row=1, max_row=max_rows, max_col=min(sheet.max_column, 16)):
        values = []
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value).replace("\n", " ").strip()
            if len(text) > 70:
                text = text[:70] + "…"
            values.append(f"{cell.coordinate}={text}")
        if values:
            print(" | ".join(values))
